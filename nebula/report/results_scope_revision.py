"""Local results-first table edit; no application or scientific mutation."""
from pathlib import Path
import copy, hashlib, json, zipfile
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from lxml import etree
from nebula.report.competition_reframe import replace

ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/'output/docx/Nebula_Competition_Report_20260915_Final.docx'
OUT=ROOT/'output/docx/Nebula_Competition_Report_20260915_Final_v2.docx'
EXPECTED='0c10e6dcadcc31b19d2f6aa86debc14485f7fa50ac51cfc788ec631dbd832349'
RENDER=ROOT/'tmp/academic-closeout/results-scope-final'
INTRO=('Nebula demonstrates an integrated RL-assisted equalizer-design workflow: target specification input, automated candidate selection, NGSpice verification, automatic rejection and recovery, and generated circuit files with resulting specifications. At 5 Gbps, one UI is 200 ps and Nyquist is 2.5 GHz. The prescribed architecture combines one source-degenerated CTLE with variable Rs/Cs and one DFE tap. The table summarizes demonstrated results and their verification scope. [E7]')
ROADMAP=('Qualification roadmap. Remaining work comprises complete target-range coverage, closure of signed device model-domain findings, complete selected-receiver analog/PVT verification, routed layout and measured-channel/BER validation. These extensions do not change the scope of the demonstrated results above.')
ROWS=[
['SPEC','REQUIRED TARGET','DEMONSTRATED RESULTS AND VERIFICATION SCOPE'],
['S1\nSignaling','5 Gbps NRZ; 2.5 GHz Nyquist','5 Gbps NRZ demonstrated: 200 ps UI and nominal transistor-receiver waveforms.'],
['S2\nCircuit','One-stage CTLE; variable Rs/Cs; one-tap DFE','Selected CTLE and nominal transistor DFE implemented. A derived receiver adds physical programmable Rs/Cs (Section 8). Accepted setting 352 remains fixed-Rs/Cs.'],
['S3\nTuning','3-12 dB; 1.25-2.5 GHz','Combined grid coverage increased from 3/12 to 4/12. Setting 352: 6.6147 dB at 2.1595 GHz. Complete on-demand range coverage remains an extension objective.'],
['S4\nLinearity','HD3 below -30 dBc; 100 MHz input','Selected CTLE nominal: -84.5045 dBc, below the limit. Independent receiver reference: about -54.92 dBc; selected-receiver HD3 remains open.'],
['S5\nNoise','Below 1.5 mVrms; 10 MHz-5 GHz','Selected CTLE maximum: 0.680353 mVrms, below the limit. Independent held-clock reference: 0.654051 mVrms; periodic receiver noise remains open.'],
['S6\nPower','Below 15 mW','Selected CTLE maximum: 9.80212 mW; selected receiver nominal: 9.2636 mW. Receiver result is VDD draw; external generators excluded.'],
['S7\nArea','130 nm PDK; below 0.05 mm2','SKY130 selected geometry subtotal: 0.00734015 mm2; independent reference: 0.014715766 mm2. Device geometry estimate, not routed layout area.'],
['S8\nEye','Above 100 mV and 0.4 UI','Selected CTLE + ideal DFE minima: 178.288 mV / 0.796875 UI. Separate nominal transistor receiver: 296.981 mV / 0.675 UI above 100 mV.'],
['S9\nPVT','TT/SS/FF/SF/FS; VDD +/-5%; 0-125 C','315/315 CTLE + ideal-DFE conditions pass: 45 PVT combinations x seven channel losses. Selected transistor receiver has separate nominal evidence.'],
]

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def check_structure():
    a,b=Document(SOURCE),Document(OUT)
    assert sha(SOURCE)==EXPECTED
    assert a.paragraphs[2].text==b.paragraphs[2].text
    assert len(a.tables)==len(b.tables)==11
    for i,(ta,tb) in enumerate(zip(a.tables,b.tables)):
        if i!=1:assert ta._tbl.xml==tb._tbl.xml,('unrelated table changed',i)
    assert [[c.text for c in r.cells] for r in b.tables[1].rows]==ROWS
    expected=[p.text for p in a.paragraphs];expected[8]=INTRO
    expected.insert(10,ROADMAP)
    assert [p.text for p in b.paragraphs]==expected,'Unrelated prose changed'
    with zipfile.ZipFile(SOURCE) as x,zipfile.ZipFile(OUT) as y:
        media=[n for n in x.namelist() if n.startswith('word/media/')]
        assert len(media)==12
        assert all(x.read(n)==y.read(n) for n in media)
        ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
        maths=lambda z:[etree.tostring(e) for e in etree.fromstring(z.read('word/document.xml')).xpath('//m:oMath',namespaces=ns)]
        assert maths(x)==maths(y) and len(maths(y))==3
    text=' '.join(p.text for p in b.paragraphs)
    for s in ('Six signed PMOS violations','Seven signed device-domain findings','does not establish an RL speed advantage','not established near-optimality'):
        assert s in text,s
    print('PASS: source/identity, unrelated tables/prose, requested table, 12 figures/3 equations, material boundaries')

def build():
    assert sha(SOURCE)==EXPECTED,'Source changed; preserve and review owner edits'
    if OUT.exists():
        prior=json.loads(OUT.with_suffix('.sources.json').read_text())
        assert sha(OUT)==prior['sha256'],'Owner changed output; preserve it'
    d=Document(SOURCE);t=d.tables[1]
    replace(d.paragraphs[8],INTRO)
    for row,values in zip(t.rows,ROWS):
        for cell,value in zip(row.cells,values):
            replace(cell.paragraphs[0],value)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    # Compact identifiers give the evidence column more space without shrinking type.
    widths=[Inches(.80),Inches(1.55),Inches(4.48)]
    for col,width in zip(t.columns,widths):col.width=width
    for row in t.rows:
        for cell,width in zip(row.cells,widths):cell.width=width
    anchor=d.paragraphs[10]
    p=anchor.insert_paragraph_before()
    p.style=anchor.style
    p._p.insert(0,copy.deepcopy(anchor._p.pPr))
    lead,body=ROADMAP.split('. ',1)
    for text,bold in ((lead+'. ',True),(body,False)):
        r=p.add_run(text);r.bold=bold;r.font.name='Times New Roman';r.font.size=Pt(10)
    p.paragraph_format.keep_together=True
    d.core_properties.comments='Results-first S1-S9 table; original results, circuit ownership and qualification scope retained.'
    d.save(OUT);check_structure()
    words=len(' '.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells]).split())
    record=dict(source=str(SOURCE),source_sha256=EXPECTED,output=str(OUT),sha256=sha(OUT),word_count=words,table=ROWS,introduction=INTRO,qualification_roadmap=ROADMAP,scientific_results_changed=False,identity_preserved=True)
    OUT.with_suffix('.sources.json').write_text(json.dumps(record,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in record.items() if k not in ('table','introduction','qualification_roadmap')},indent=2))

if __name__=='__main__':
    import sys
    if '--check' in sys.argv:check_structure()
    else:build()
