"""Read-only structural checks for the final presentation artifacts."""
from pathlib import Path
import hashlib,json,re,zipfile
from docx import Document
from pypdf import PdfReader
from lxml import etree
from PIL import Image,ImageChops

ROOT=Path(__file__).resolve().parents[2]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    meta=json.loads((ROOT/'output/docx/Nebula_Current_Presentation_20260915_v2_sources.json').read_text())
    for item in ('report','demo'):assert sha(Path(meta[item]['path']))==meta[item]['sha256']
    source=ROOT/'output/docx/Nebula_Academic_Final_Report_20260915_v2.docx'
    assert sha(source)==meta['report']['source_sha256']
    old,new=Document(source),Document(meta['report']['path'])
    assert [[c.text for r in t.rows for c in r.cells] for t in old.tables]==[[c.text for r in t.rows for c in r.cells] for t in new.tables]
    prefixes=meta['report']['changed_paragraph_prefixes']
    assert len(old.paragraphs)==len(new.paragraphs)
    for a,b in zip(old.paragraphs,new.paragraphs):
        if a.text!=b.text:assert any(a.text.startswith(prefix) for prefix in prefixes)
    with zipfile.ZipFile(meta['report']['path']) as z:
        doc=etree.fromstring(z.read('word/document.xml'))
        assert len(doc.xpath('//m:oMath',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}))==3
    assert len(new.inline_shapes)==12
    report_dir=ROOT/'tmp/academic-closeout/entry163-report-v4'
    demo_dir=ROOT/'tmp/academic-closeout/entry163-demo-v2'
    pages={}
    for label,directory in [('report',report_dir),('demo',demo_dir)]:
        pdf=next(directory.glob('*.pdf'));pages[label]=len(PdfReader(pdf).pages)
        assert all((directory/f'page-{n}.png').is_file() for n in range(1,pages[label]+1))
    assert pages==dict(report=12,demo=4)
    report_text=[p.extract_text() for p in PdfReader(next(report_dir.glob('*.pdf'))).pages]
    assert 'Programmable receiver extension' in report_text[7]
    assert 'Run files / audit.' in '\n'.join(report_text[9:11])
    assert 'port 8765 is the legacy server' in ' '.join(' '.join(report_text).split())
    identical=[]
    for n in range(1,13):
        a=Image.open(ROOT/f'tmp/academic-closeout/entry163-report-render/page-{n}.png').convert('RGB');b=Image.open(report_dir/f'page-{n}.png').convert('RGB')
        if a.size==b.size and ImageChops.difference(a,b).getbbox() is None:identical.append(n)
    from nebula.programmable_option import for_parent
    data=for_parent(ROOT/'nebula/product_demo/submission_runs_20260915_v2/4608cf1c525f4e59b2d5226059b0ce5d')
    assert data['correct_bits']==64 and not data['full_receiver_verified']
    for name in ('entry163-rehearsal.json','entry163-supplemental.json'):json.loads((ROOT/'tmp/academic-closeout'/name).read_text())
    print(json.dumps(dict(status='PASS',pages=pages,unchanged_pixels_from_reviewed_first_render=identical,report_sha256=meta['report']['sha256'],demo_sha256=meta['demo']['sha256'],scientific_tables_unchanged=True,core_spoken_words=meta['demo']['core_spoken_words']),indent=2))
if __name__=='__main__':main()
