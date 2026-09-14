"""Build the final own-voice narration as Markdown and Word; no scientific runs."""
from __future__ import annotations
import hashlib
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
BASENAME = 'Nebula_Final_Demo_Script_20260915'
OUT = ROOT / 'output/docx' / (BASENAME + '.docx')
MD = ROOT / 'nebula/SUBMISSION_FINAL_DEMO_SCRIPT_20260915.md'
REPORT = 'Nebula_Final_Submission_12p_20260915.pdf'
SCENES = [
('The design problem', '0:00 to 0:35',
 'Show final report page 1 for the opening, then Design Explorer with saved physical run ab981d69dac840e9820155d7e46e14dc selected. Keep its result and CTLE visible.', [
 'Hello. This is Nebula, our RL-assisted Python framework for designing a five-gigabit-per-second NRZ equalizer using SKY130 and ngspice.',
 'The challenge is to accept electrical targets, find useful circuit settings, verify them and return a schematic with its resulting specifications. I will show the automated design flow, a completed physical CTLE, and a separate transistor CTLE-plus-DFE checkpoint.']),
('How the design flow works', '0:35 to 1:20',
 'Open New target, point to peaking and peak-frequency inputs, then close without generating. Briefly show final report pages 3 and 4, then return to Explorer.', [
 'The trained policy proposes attenuation, resistance and capacitance codes within a characterized library of five hundred twelve settings. It uses the request, current code and measured eye history to choose its next move.',
 'A deterministic verifier checks the electrical and link constraints. Safe candidates are ranked for target accuracy, and a classical fallback can recover a result when the policy misses one. Physical generation then measures the selected fixed CTLE. This separates learned proposals from the measurements that justify acceptance.']),
('Inspect a completed circuit and its provenance', '1:20 to 2:35',
 'Select Response, then Sizing beside the CTLE canvas. Expand the exact generated CTLE drawing. Show final report page 10 for the SEPARATE saved adaptive-run receipt, then return to the 3 dB physical result.', [
 'This is a saved completed physical run. Its request was three decibels at one point nine gigahertz. It measures three point nine one decibels at two point one three one gigahertz, within this path\'s stated tolerances.',
 'The drawing includes the differential pair and both tail current sources. Source degeneration shapes the response, while a physical reference supplies the bias. The inspector shows values and specifications beside the circuit; the exported drawing and deck preserve the exact implementation.',
 'The report also traces a separate saved nine-decibel adaptive run. It records two thousand four hundred six visits, including fixed starts and repeats. Forty-nine conditions needed safety recovery. In the example shown, the policy never visited the accepted fallback code. The receipt keeps that distinction explicit, and its recorded timer is not a complete end-to-end speed measurement.']),
('Check corners and the modeled eye', '2:35 to 3:30',
 'Return to the saved 3 dB PHYSICAL run. Click Inspect all PVT corners, select a cell and point to the loss selector. Open Compare circuits and choose that physical run as Circuit A. Switch CTLE output to Ideal 1-tap DFE; briefly show Margin envelope.', [
 'For this fixed physical CTLE, three hundred fifteen model conditions pass: forty-five PVT corners across seven constructed channel losses. Each grid shows one loss slice, with the overall count kept separate.',
 'The eye overlay uses this circuit\'s saved transistor AC response. Ideal one-tap feedback subtracts the first postcursor contribution. The recorded conservative margins are about three hundred forty-two millivolts and zero point eight five nine UI.',
 'Those dimensions come from cursor analysis, separately from the finite waveform overlay. This is noiseless modeled link evidence; it does not establish transistor transient behavior or BER.']),
('Show the separate transistor checkpoint', '3:30 to 4:40',
 'Return to Explorer. Expand the separate checkpoint labeled 9 dB / 1.9 GHz. Show its CTLE plus DFE circuit, Rs/Cs controls and saved eye/PVT evidence. Final report pages 8 and 9 carry the corresponding engineering and results.', [
 'This independent checkpoint implements the physical feedback path: configurable resistance and capacitance, the CTLE, a CML summer, decision-memory latches and a feedback current DAC. The calibrated deck contains seventy-three SKY130 device instances.',
 'It was developed through bounded circuit experiments alongside the framework, rather than autonomously discovered by the frozen policy. Loading moved the frequency response, so the actual loaded controls were calibrated and independently checked.',
 'Nominal peaking is about eight point seven four decibels near one point nine zero five gigahertz. Nominal held-state noise and clocked distortion meet their declared limits.',
 'The same controls, tap code and clock phase pass all forty-five sampled link PVT points without retuning. The minimum sampled eye is about one hundred thirteen millivolts. This finite noiseless check and the separate analog measurements must retain their stated coverage.']),
('Deliverables and remaining work', '4:40 to 5:40',
 'Show Run files and an exact deck or schematic. Open final report page 2 for the requirement matrix and page 12 for the artifact guide. Finish on the selected physical CTLE in Explorer.', [
 'The deliverable connects target input, learned proposals, deterministic verification and ngspice to reviewable circuit outputs. The report explains both successful designs and the failures that shaped the engineering decisions.',
 'Across five frozen seeds, PPO improves normalized compliant eye quality over a fixed start and imitation. Local random search achieves higher compliance, and near-optimality is not established. These are useful measured trade-offs, not a claim that RL wins every comparison.',
 'Full target-range coverage, simultaneous all-specification PVT compliance, periodic receiver noise and routed layout remain open. VDD power excludes external clock and control sources. Our contribution is a working RL-assisted framework with traceable physical checkpoints that an analog designer can inspect and reproduce. Thank you.'])]


def build() -> Path:
    doc=Document()
    # Remove inherited template borders and theme fonts for portable rendering.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for fonts in doc.styles.element.iter(qn('w:rFonts')):
        for key in list(fonts.attrib):
            if key.endswith('Theme') or key.endswith('theme'):
                del fonts.attrib[key]
    sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.60);sec.bottom_margin=Inches(.60)
    sec.left_margin=sec.right_margin=Inches(.75)
    normal=doc.styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(12)
    normal.font.color.rgb=RGBColor.from_string('000000')
    normal.paragraph_format.line_spacing=1.07
    normal.paragraph_format.space_after=Pt(7)
    for name,size in [('Title',22),('Heading 1',15),('Heading 2',12)]:
        st=doc.styles[name];st.font.name='Aptos';st.font.size=Pt(size)
        st.font.color.rgb=RGBColor.from_string('000000')
        st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.keep_with_next=True
        if name!='Title':st.font.bold=True
    doc.add_paragraph('Nebula final demo narration','Title')
    p=doc.add_paragraph('Own voice walkthrough  |  5 to 6 minutes  |  15 September 2026')
    p.runs[0].font.size=Pt(10)
    p=doc.add_paragraph('Presenter: ____________________    Team: ____________________')
    p.runs[0].font.size=Pt(10)
    p=doc.add_paragraph('Read the black narration aloud. Blue screen cues are for you. Use saved completed evidence and pause briefly after clicks.')
    p.runs[0].font.size=Pt(10.5)
    md=['# Nebula final demo narration','',
        'Own voice walkthrough. Target duration: 5-6 minutes; the cue schedule is an estimate, not a recorded duration.',
        '', 'Presenter: ____________________    Team: ____________________', '',
        'Use `http://127.0.0.1:8765` and the final 12-page report `output/pdf/'+REPORT+'`. Read narration only. Saved evidence is not a live optimization.','']
    for i,(title,timing,screen,paras) in enumerate(SCENES):
        if i in (2,4):doc.add_page_break()
        doc.add_paragraph(title,'Heading 1')
        p=doc.add_paragraph(timing);p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(4);p.runs[0].bold=True;p.runs[0].font.size=Pt(10)
        p=doc.add_paragraph('Screen  '+screen)
        p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(8)
        for r in p.runs:r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string('24536B')
        for text in paras:
            p=doc.add_paragraph(text);p.paragraph_format.widow_control=True
        md += ['## '+title, '', '**'+timing+'**', '', '**Screen:** '+screen,'', *sum(([p,''] for p in paras),[])]
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Nebula narration  |  ').font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title='Nebula final demo narration'
    doc.core_properties.subject='Own voice walkthrough of the RL-assisted equalizer design framework'
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    OUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUT)
    MD.write_text('\n'.join(md)+'\n',encoding='utf-8',newline='\n')
    words=sum(len(re.findall(r"\b[\w'-]+\b",p)) for s in SCENES for p in s[3])
    check={'spoken_word_count':words,'scenes':len(SCENES),'target_minutes':'5-6',
           'pace_words_per_minute_at_5min':words/5,'pace_words_per_minute_at_6min':words/6,
           'docx_path':OUT.relative_to(ROOT).as_posix(),
           'docx_sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),
           'source':Path(__file__).relative_to(ROOT).as_posix(),
           'source_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
           'report':REPORT,'report_pages_used':[1,2,3,4,8,9,10,12],
           'visual_review':'PENDING','video_reviewed':False}
    OUT.with_suffix('.content.json').write_text(json.dumps(check,indent=2)+'\n',encoding='utf-8')
    return OUT


if __name__=='__main__':
    print(build())
