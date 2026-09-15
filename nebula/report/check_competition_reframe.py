"""Focused editorial integrity checks; no scientific evaluation or mutation."""
from pathlib import Path
import hashlib,json,re,unittest,zipfile
from docx import Document
from lxml import etree
from pypdf import PdfReader
from PIL import Image,ImageChops
from nebula.report import competition_reframe as C

class CompetitionReportChecks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.old=Document(C.SOURCE);cls.new=Document(C.OUT)
        cls.record=json.loads((C.ROOT/'output/docx/Nebula_Competition_Report_20260915_Final_sources.json').read_text())

    def test_source_and_identity_preserved(self):
        self.assertEqual(C.sha(C.SOURCE),C.EXPECTED_SOURCE)
        self.assertEqual(C.sha(C.OUT),self.record['sha256'])
        self.assertEqual(self.old.paragraphs[2].text,self.new.paragraphs[2].text)
        self.assertEqual([[c.text for c in r.cells] for r in self.old.tables[0].rows],[[c.text for c in r.cells] for r in self.new.tables[0].rows])

    def test_only_two_declared_table_cells_change(self):
        changed=[]
        for ti,(a,b) in enumerate(zip(self.old.tables,self.new.tables)):
            for ri,(ar,br) in enumerate(zip(a.rows,b.rows)):
                for ci,(ac,bc) in enumerate(zip(ar.cells,br.cells)):
                    if ac.text!=bc.text:changed.append((ti,ri,ci))
        self.assertEqual(changed,[(1,3,0),(1,3,2)])
        row=self.new.tables[1].rows[3]
        self.assertIn('4/12',row.cells[2].text)
        self.assertIn('Full on-demand range not met',row.cells[2].text)
        self.assertNotIn('pass',row.cells[0].text.lower())

    def test_figures_and_native_equations_preserved(self):
        self.assertEqual(len(self.new.inline_shapes),12)
        with zipfile.ZipFile(C.SOURCE) as a,zipfile.ZipFile(C.OUT) as b:
            for name in [n for n in a.namelist() if n.startswith('word/media/')]:self.assertEqual(a.read(name),b.read(name))
            ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
            maths=lambda z:[etree.tostring(e) for e in etree.fromstring(z.read('word/document.xml')).xpath('//m:oMath',namespaces=ns)]
            self.assertEqual(maths(a),maths(b));self.assertEqual(len(maths(b)),3)

    def test_material_boundaries_and_criteria_remain(self):
        text=' '.join(p.text for p in self.new.paragraphs)
        for phrase in ('Deliverable coverage.','Innovation.','Thought process.','ideal behavioral one-tap DFE',
            'Six signed PMOS violations','Seven signed device-domain findings','rejected noise instrument',
            'full receiver verification remains future work','does not establish an RL speed advantage',
            'not established near-optimality','Eyes are noiseless ISI openings, not BER contours',
            'live provider is unavailable','Setting 352 remains the primary fixed output'):
            self.assertIn(phrase,text)

    def test_twelve_rendered_pages_and_figure_order(self):
        folder=C.ROOT/'tmp/academic-closeout/competition-report-final'
        reader=PdfReader(folder/(C.OUT.stem+'.pdf'));self.assertEqual(len(reader.pages),12)
        self.assertTrue(all((folder/f'page-{n}.png').is_file() for n in range(1,13)))
        figures=[int(re.match(r'Figure (\d+)\.',p.text)[1]) for p in self.new.paragraphs if re.match(r'Figure (\d+)\.',p.text)]
        self.assertEqual(figures,list(range(1,13)))

if __name__=='__main__':unittest.main()
