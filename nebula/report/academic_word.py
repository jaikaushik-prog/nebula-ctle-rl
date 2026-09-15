"""Plain academic Word report, authored from the verified current report content."""
from pathlib import Path
import hashlib
import html
import json
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'output/docx'
REPORT=OUT/'Nebula_Academic_Final_Report_20260915.docx'
TITLES=['Abstract','Problem statement and requirements','Framework architecture and target coverage',
 'Reinforcement learning formulation and evaluation','Selected physical circuit implementation',
 'Peaking attenuation and programmable controls','Device to link modeling',
 'Generated transistor receiver verification','Independent reference verification',
 'Automatic recovery and complete design cost','Software interface and operation','Conclusions and reproducibility']

def text(s):
    s=re.sub(r'<br\s*/?>','; ',str(s));s=html.unescape(re.sub('<[^>]+>','',s))
    s=s.replace('companion v4 source manifest','companion PDF source manifest')
    s=s.replace(';;',';')
    s=s.replace('12-page report','technical report').replace('separate v4 outputs','separate outputs')
    s=s.replace('p.8','Section 8').replace('pp.','Sections ').replace('p.12','Section 12')
    return s

def base(title):
    d=Document();s=d.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69)
    s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.72)
    s.footer_distance=Inches(.28)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption'):
        st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Times New Roman')
    n=d.styles['Normal'];n.font.size=Pt(12);n.paragraph_format.line_spacing=1.08
    n.paragraph_format.space_after=Pt(5);n.paragraph_format.widow_control=True
    for name,size in [('Title',17),('Heading 1',13),('Heading 2',11.5),('Caption',9)]:
        d.styles[name].font.size=Pt(size)
    for name in ('Heading 1','Heading 2'):
        st=d.styles[name];st.font.bold=True;st.paragraph_format.space_before=Pt(10)
        st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
    d.styles['Title'].paragraph_format.space_after=Pt(8)
    f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    for element in (d.styles.element,d.element):
        for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
        for fonts in element.iter(qn('w:rFonts')):
            for attr in list(fonts.attrib):
                if attr.split('}')[-1].endswith('Theme'):del fonts.attrib[attr]
            for attr in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+attr),'Times New Roman')
    d.styles['Caption'].font.bold=False;d.styles['Caption'].font.italic=True
    d.add_paragraph(title,'Title');d.core_properties.title=title;d.core_properties.author=''
    return d

def para(d,s,style=None):
    p=d.add_paragraph(text(s),style);return p

def table(d,heads,rows,widths=None):
    t=d.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or ([1.15,1.55,4.13] if len(heads)==3 else [6.83/len(heads)]*len(heads))
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
    for i,values in enumerate([heads,*rows]):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(c,value) in enumerate(zip(cells,values)):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            c.text=text(value);pr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for edge,v in [('top','65'),('bottom','65'),('left','85'),('right','85')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),v);e.set(qn('w:type'),'dxa');m.append(e)
            pr.append(m)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');pr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
                for r in p.runs:r.font.size=Pt(10);r.bold=i==0
            cant=OxmlElement('w:cantSplit');c._tc.getparent().get_or_add_trPr().append(cant)
    d.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def figure(d,b,number):
    from PIL import Image
    h={'architecture':1.3,'learning':1.95,'selected_ctle':3.75,'submission_attenuator':1.45,
       'submission_ac':1.7,'submission_eyes':1.6,'submission_dfe':1.6,'receiver':1.1,
       'generated_aperture':1.65,'pvt':1.75,'desktop':2.15}[b['name']]
    path=ROOT/'nebula/report/academic_final_assets/desktop.png' if b['name']=='desktop' else Path(b['path']);w,ih=Image.open(path).size;scale=min(6.7/w,h/ih)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(w*scale),height=Inches(ih*scale))
    d.inline_shapes[-1]._inline.docPr.set('descr',text(b['caption']))
    p=para(d,f'Figure {number}. '+b['caption'],'Caption');p.paragraph_format.space_after=Pt(7)

def build_report():
    data=json.loads((ROOT/'nebula/report/academic_final_assets/report_content.json').read_text())
    d=base('Nebula Reinforcement Learning Based Equalizer Design')
    para(d,'Final technical report | PCIe Gen2 5 Gbps NRZ | SKY130 130 nm | 15 September 2026')
    para(d,'Team ____________________    Institution ____________________')
    table(d,['Member name','Contact number','Email address'],[['','',''] for _ in range(3)],[2.4,1.6,2.83])
    figures=0
    narrative_heads={'JUDGING CRITERION','EVIDENCE','PART OF THE MDP','DESIGN STEP','WORKFLOW','CIRCUIT / RECORD'}
    for section in data['sections']:
        n=section['number'];d.add_heading(f'{n} {TITLES[n-1]}',1)
        # Academic prose below supplies context; omit the designer-style subtitle.
        for b in section['blocks']:
            kind=b['kind']
            if kind=='identity':continue
            if kind=='figure':
                if n==1:continue
                figures+=1;figure(d,b,figures);continue
            if kind=='heading':
                title=text(b['text'])
                if title.startswith('Physical target evidence:'):title='Tested target coverage from three to four of twelve'
                elif title=='From separate geometries to physical variable Rs/Cs':title='Physical variable resistance and capacitance'
                elif title=='What the 1-tap DFE contributes':title='Contribution of one tap decision feedback'
                d.add_heading(re.sub(r'[^\w\s]',' ',title).strip(),2);continue
            if kind=='table':
                heads=b['headings'];rows=b['rows']
                if n==1:continue
                if heads[0] in narrative_heads:
                    for row in rows:
                        content=' '.join(text(v) for v in row[1:])
                        content=re.sub(r'pp\.[0-9, -]+;?','',content)
                        p=para(d,'');p.add_run(text(row[0])+'. ').bold=True;p.add_run(content)
                elif heads[0]=='ADDITIONAL EVIDENCE':
                    for row in rows:para(d,text(row[0])+'. '+text(row[1]))
                else:
                    widths=([1.1,1.55,4.18] if n==2 else [1.9,1.65,3.28] if len(heads)==3 else [2.3,4.53] if len(heads)==2 else [2.65,1.39,1.39,1.4])
                    table(d,heads,rows,widths)
                continue
            s=b['text']
            if n==1 and (s.startswith('<b>Specifications') or s.startswith('<b>Contribution.')):continue
            if n==12 and (s.startswith('<b>Methods.') or s.startswith('<b>Exact paths')):continue
            if s.startswith('<b>First-order reasoning.'):
                para(d,'First-order design equations include source degeneration and body effect. Rs denotes the full resistance between the two sources. The following approximations support interpretation; final values come from SPICE, which includes finite output resistance and parasitics. [E7,E11]')
                def mr(value):
                    run=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;run.append(t);return run
                def sub(value,index):
                    node=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(value));s=OxmlElement('m:sub');s.append(mr(index));node.extend([e,s]);return node
                def frac(top,bottom):
                    node=OxmlElement('m:f');a=OxmlElement('m:num');b=OxmlElement('m:den');a.extend(top);b.extend(bottom);node.extend([a,b]);return node
                equations=[
                    [sub('f','z'),mr(' = '),frac([mr('1')],[mr('2\u03c0 '),sub('R','s'),sub('C','s')])],
                    [mr('k = 1 + '),frac([mr('('),sub('g','m'),mr(' + '),sub('g','mb'),mr(')'),sub('R','s')],[mr('2')])],
                    [sub('A','dc'),mr(' \u2248 '),frac([sub('g','m'),sub('R','L')],[mr('k')])]]
                for nodes in equations:
                    p=d.add_paragraph();p.paragraph_format.space_after=Pt(3);math=OxmlElement('m:oMath');math.extend(nodes);p._p.append(math)
                continue
            para(d,s)
        if n==3:
            para(d,'The catalogue reduces the online decision to a bounded search over characterized choices. It does not remove the offline characterization cost or prove superiority over an unrestricted MOS, resistor, capacitor and inductor sweep. The continuous-Cs extension is separately registered local refinement, not an expansion discovered by the trained policy.')
        if n==9:
            rows=[]
            for t in data['diagnostics']['tuning']['targets']:
                if t['nominal_ac_match']:rows.append([f"{t['target_boost_db']:g} dB / {t['target_frequency_hz']/1e9:g} GHz",f"{min(t['measured_boost_db']):.3f}-{max(t['measured_boost_db']):.3f} dB",f"{min(t['measured_peak_ghz']):.3f}-{max(t['measured_peak_ghz']):.3f} GHz"])
            table(d,['Loaded target','Measured boost','Measured peak'],rows,[2.3,2.1,2.43])
        if n==10:
            para(d,'The rejected 288 candidate is not a demonstrated closed transistor eye. At SF, 0.95 VDD and 0 C with the 3 dB constructed channel, the requested output swing exceeded the linear-model guard. The pipeline therefore refused the bridge calculation at that condition and retained the failure before trying setting 352. This distinction prevents an invalid model output from becoming a plausible eye measurement.')
        if n==5:
            para(d,'The simulator integration parses primitive operating-point, AC, noise and transient outputs. Simulator exit status alone is insufficient because ngspice can report some failures as warnings. The framework rejects malformed or nonfinite data and invalid instruments. HD3 uses transient harmonic extraction rather than treating a zero BSIM4 distortion-analysis output as a valid linearity result. Noise integration is reported as RMS volts; it is not square-rooted a second time.')
        if n==11:
            para(d,'The interface separates request entry, candidate acceptance and evidence inspection. Reading a sentence only parses the requested values; it does not start a simulation. Generate and verify starts a fresh protected workflow. Selecting a saved run replays its existing evidence. Downloads expose recorded files without remeasuring the circuit.')
    d.add_heading('References and evidence records',2)
    for s in [
      'E1-E2. Independent calibrated transistor reference and 45-point link PVT. nebula/product_audits/entry143_dfe_calibrated_verification_20260910 and entry144_dfe_calibrated_pvt_20260910. Exact locations and hashes are bound by the preserved PDF source manifest.',
      'E3-E5. Frozen FINAL policy evaluation, exposed attribution controls and exhaustive bank-oracle comparison. nebula/experiments/shielded_policy_final_results.json; nebula/product_audits/entry113_attribution_20260907; entry115_exhaustive_benchmark_20260908.',
      'E6-E10. Historical fixed CTLE, specification contract, adaptive-bank receipt and invalid analog-PVT pilot. See the companion source manifest for exact ownership and archived paths.',
      'E11-E14. Recorded recovery, matched workflow benchmark, continuous-Cs study and nominal selected receiver. nebula/product_audits/submission_recovery_20260915; continuous_cs_coverage_20260915; generated_receiver_6db_2p1ghz_20260915.',
      'E15-E16. Bounded loaded tuning and read-only cost/model-domain review. nebula/product_audits/entry158_loaded_tuning_20260915; nebula/DEADLINE_ENGINEERING_RESULTS_20260915.md; nebula/deadline_evidence.py.',
      'SkyWater PDK documentation. https://skywater-pdk.readthedocs.io/en/main/ . ngspice documentation. https://ngspice.sourceforge.io/docs.html .',
      'Schulman J et al. Proximal Policy Optimization Algorithms. 2017. https://arxiv.org/abs/1707.06347 .',
      'Guo et al. Balancing Speed and Accuracy for Robust Analog-Mixed Signal Circuit Design using Closed-Loop Reinforcement Learning with Ensemble Neural Network Surrogates. ISCAS 2026. Methodological reference supplied for the competition; its published results are not Nebula measurements.'
    ]:para(d,s)
    OUT.mkdir(exist_ok=True);d.save(REPORT)
    return dict(path=str(REPORT),figures=figures,sha256=hashlib.sha256(REPORT.read_bytes()).hexdigest())

if __name__=='__main__':print(json.dumps(build_report(),indent=2))
