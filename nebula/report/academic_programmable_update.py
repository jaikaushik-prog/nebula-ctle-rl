"""Preserve the current academic DOCX and add the measured programmable variant."""
from pathlib import Path
import hashlib,json,copy,re
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from nebula.report.academic_word import table,para
from nebula import programmable_option as P

ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/'output/docx/Nebula_Academic_Final_Report_20260915.docx'
OUT=ROOT/'output/docx/Nebula_Academic_Final_Report_20260915_v2.docx'

def replace(p,value):
    # Retain paragraph/run typography, but not stale textual content.
    style=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear();r=p.add_run(value)
    if style is not None:r._r.insert(0,style)

def build():
    data=P._review();d=Document(SOURCE);original_sha=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    paragraphs=list(d.paragraphs)
    def find(prefix):return next(p for p in paragraphs if p.text.startswith(prefix))
    body_format=copy.deepcopy(find('Saved transistor AC').runs[-1]._r.rPr)
    def body_style(p):
        for run in p.runs:
            if run._r.rPr is not None:run._r.remove(run._r.rPr)
            if body_format is not None:run._r.insert(0,copy.deepcopy(body_format))
        return p
    p=find('Nebula turns target specifications');replace(p,p.text+' A programmable Rs/Cs receiver has also been implemented as an experimental option. [E17]')
    cell=d.tables[1].rows[2].cells[2]
    replace(cell.paragraphs[0],'Selected 352: fixed Rs/Cs; transistor DFE nominal. A derived programmable receiver now adds physical Rs/Cs controls (Section 8); full receiver verification remains future work.')
    anchor=find('Independent 9 dB reference.')
    p=body_style(anchor.insert_paragraph_before('Programmable selected receiver. A separate variant of setting 352 connects physical Rs/Cs controls to the transistor DFE. Nominal measured controls and signal results are recorded separately. [E17]'))
    p=find('Figure 8.');replace(p,p.text+' The selected-derived programmable option below uses this receiver topology with a replacement Rs/Cs network. [E17]')
    for old,new in [('Figure 11.','Figure 12.'),('Figure 10.','Figure 11.')]:
        p=next(p for p in paragraphs if p.text.startswith(old));replace(p,new+p.text[len(old):])
    # Insert a compact new subsection within receiver Section 8, before Section 9.
    anchor=find('9 Independent reference verification')._p
    def insert(p):
        if p.style.name=='Normal' and p.text:body_style(p)
        anchor.addprevious(p._p);return p
    insert(d.add_heading('Programmable receiver extension',2))
    insert(para(d,'We integrated a controlled resistor branch and two SKY130 varactors into a receiver derived from setting 352. The amplifier, attenuation, bias and transistor DFE remain unchanged. External R and C controls of 1.314 V and 0.378 V select the measured operating point. This gives the framework a physical tuning option connected to decision feedback, rather than only a standalone tuning circuit. [E17]'))
    t=table(d,['NOMINAL MEASUREMENT','RESULT','CONDITIONS'],[
      ['Peaking / peak frequency','5.615-5.618 dB; 2.068-2.071 GHz','Both held-clock states; 6 dB / 2.1 GHz request within existing +/-0.5 dB / +/-100 MHz internal tolerances.'],
      ['Decisions / sampled eye','64/64; 325.629 mV','TT / 1.8 V / 27 C; constructed 7.5 dB channel; clock phase 1 UI, tap code 2.'],
      ['Aperture / VDD power','0.720 UI; 9.2643 mW','Scan-limited positive aperture; 0.680 UI above 100 mV. External source power excluded.']],[1.7,2.0,3.13])
    anchor.addprevious(t._tbl)
    # Move the spacing paragraph created by the shared table helper with its table.
    spacer=d.paragraphs[-1];anchor.addprevious(spacer._p)
    p=insert(d.add_paragraph());p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(P.artifact('response-eye')),width=Inches(6.7))
    p._p.xpath('.//wp:docPr')[0].set('descr','Measured loaded CTLE response at three Cs controls and nominal programmable transistor receiver eye')
    insert(para(d,'Figure 10. Loaded response at three fixed Cs controls, with R control held at 1.314 V, and the chosen transistor receiver eye. These saved measurements are not a runtime retuning demonstration. [E17]','Caption'))
    insert(para(d,'Nominal simulation results are shown; full receiver verification remains future work. Signed model-domain findings, a rejected noise instrument and outstanding HD3/PVT checks are retained in [E17], together with all eight simulator calls. Control calibration is deterministic, not RL. The accepted setting-352 circuit remains the primary submission; grid coverage stays 4/12.'))
    p=find('Run files / audit.');replace(p,p.text+' The expandable programmable receiver panel adds its own measured deck, drawing and plots; matching future 6 dB / 2.1 GHz exports include this saved option separately without fresh verification.')
    p=find('Run files / audit.');body_style(p)
    content=p.text;p.clear();p.add_run('Run files / audit. ').bold=True;p.add_run(content[len('Run files / audit. '):])
    for run in p.runs:run.font.size=find('Saved transistor AC').runs[-1].font.size
    p=find('Three evidence branches');replace(p,'Four evidence branches with separate measurement scope')
    anchor=find('Independent 9 dB / 1.9 GHz;')
    body_style(anchor.insert_paragraph_before('Programmable selected receiver; E17 / Sections 3,8,11. Separate experimental deck and exact device drawing, with fixed external controls at 0.73/0.21 VDD. Nominal results do not inherit the accepted CTLE gate or the independent reference PVT.'))
    anchor=find('SkyWater PDK documentation.')
    body_style(anchor.insert_paragraph_before('E17. Selected-derived programmable receiver, raw-data recomputation and all retained attempts. nebula/product_audits/entry160_programmable_receiver_review_20260915/review.json; exact artifact paths and SHA-256 hashes are in this record. Integration: nebula/programmable_receiver.py and nebula/programmable_option.py.'))
    # Keep the two-arm comparison together instead of splitting one row per page.
    benchmark=next(t for t in d.tables if t.cell(0,0).text=='ARM')
    for row in benchmark.rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    find('References and evidence records').paragraph_format.page_break_before=True
    d.core_properties.comments='Updated with the separately measured programmable receiver option on 15 September 2026.'
    d.save(OUT)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==original_sha
    words=' '.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells]).split()
    record=dict(path=str(OUT),source=str(SOURCE),source_sha256=original_sha,sha256=hashlib.sha256(OUT.read_bytes()).hexdigest(),word_count=len(words),figures=len(d.inline_shapes),review_sha256=P.REVIEW_SHA)
    (ROOT/'output/docx/Nebula_Academic_Final_Report_20260915_v2_sources.json').write_text(json.dumps(record,indent=2),encoding='utf-8')
    print(json.dumps(record,indent=2))

if __name__=='__main__':build()
