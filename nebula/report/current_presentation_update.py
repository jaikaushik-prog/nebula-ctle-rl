"""Targeted academic report edit and new narration from current verified evidence."""
from pathlib import Path
import copy,hashlib,json,re
from docx import Document
from docx.shared import Pt
from nebula.report.academic_word import base

ROOT=Path(__file__).resolve().parents[2]
SOURCE=ROOT/'output/docx/Nebula_Academic_Final_Report_20260915_v2.docx'
REPORT=ROOT/'output/docx/Nebula_Academic_Final_Report_20260915_v4.docx'
SCRIPT=ROOT/'nebula/CURRENT_APP_DEMO_20260915.md'
DEMO=ROOT/'output/docx/Nebula_Current_App_Demo_20260915_v2.docx'
SCREEN=ROOT/'tmp/academic-closeout/entry163-report-workbench.png'

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def replace(p,text):
    fmt=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear();r=p.add_run(text)
    if fmt is not None:r._r.insert(0,fmt)

def build():
    before=sha(SOURCE);d=Document(SOURCE)
    changes={
      'Figure 12.':'Figure 12. Current 1440x1000 Design explorer: accepted setting 352, saved 288-to-352 recovery and fixed-CTLE scope. Drawing and receiver evidence open in separate full-width panels; the programmable prototype is under Run files. This is a saved replay, not a new design run. [E9,E11,E17]',
      'Modeled DFE / new receiver.':'Modeled DFE / new receiver. The 315-condition gate uses ideal behavioral DFE outside design.cir. Open Selected CTLE + transistor DFE for the exact-hash-matched nominal check. The generated drawing and independent reference use separate full-width panels; no receiver inherits another circuit\'s verification.',
      'Run files / audit.':'Run files / audit. Open Programmable receiver prototype for its nominal measurements, exact drawing, plots and eight artifact links. Saved data loads only when opened and stays separate from setting 352. The ownership index covers the primary CTLE, matching nominal receiver and independent reference. Prototype files are downloaded separately; older archives are unchanged.',
      'Saved-run launch:':'Saved-run launch: py -3.13 -m nebula.web.recovery_server --no-browser. Open http://127.0.0.1:8766/#results; port 8765 is the legacy server. Default loading uses saved submission runs and launches no SPICE. Temporary evidence errors offer a read-only retry. Preflight: py -3.13 -m nebula.submission_preflight.',
    }
    for prefix,value in changes.items():
        matches=[p for p in d.paragraphs if p.text.startswith(prefix)];assert len(matches)==1
        p=matches[0]
        replace(p,value)
        if prefix in ('Modeled DFE / new receiver.','Run files / audit.'):
            # Preserve the original bold lead-in without making the entire paragraph bold.
            fmt=copy.deepcopy(p.runs[0]._r.rPr)
            p.clear();lead=p.add_run(prefix+' ');body=p.add_run(value[len(prefix)+1:])
            for run in (lead,body):
                if fmt is not None:run._r.insert(0,copy.deepcopy(fmt))
            lead.bold=True;body.bold=False
            p.paragraph_format.keep_together=True
    image=d.inline_shapes[-1];rid=image._inline.xpath('.//a:blip')[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
    d.part.related_parts[rid]._blob=SCREEN.read_bytes()
    image._inline.docPr.set('descr','Current Nebula Design explorer showing the accepted fixed CTLE and recorded candidate rejection and recovery')
    d.core_properties.comments='Navigation and interface image aligned on 15 September 2026; scientific results unchanged.'
    d.save(REPORT);assert sha(SOURCE)==before
    old=Document(SOURCE)
    assert [[c.text for c in r.cells] for r in d.tables[0].rows]==[[c.text for c in r.cells] for r in old.tables[0].rows]
    assert len(d.inline_shapes)==12
    for p,q in zip(old.paragraphs,d.paragraphs):
        if p.text!=q.text:assert any(p.text.startswith(x) for x in changes)
    demo=base('Nebula Current Application Demonstration')
    demo.styles['Normal'].font.size=Pt(11)
    demo.styles['Normal'].paragraph_format.line_spacing=1.05
    content=SCRIPT.read_text(encoding='utf-8').splitlines();scene=0;spoken=0;in_core=True
    for line in content[1:]:
        if not line.strip():continue
        if line.startswith('## '):
            heading=line[3:]
            timed=bool(re.match(r'\d\d:',heading))
            if timed:
                scene+=1
                if scene in (4,7):demo.add_page_break()
                times,title=re.match(r'(\d{2}:\d{2} to \d{2}:\d{2}) (.+)',heading).groups()
                demo.add_heading(title,2);p=demo.add_paragraph(times);p.paragraph_format.space_after=Pt(4)
                for r in p.runs:r.bold=True;r.font.size=Pt(10)
            else:
                if in_core:demo.add_page_break()
                in_core=False;demo.add_heading(heading,2)
        elif line.startswith('Action: '):
            p=demo.add_paragraph();p.add_run('On screen. ').bold=True;p.add_run(line[8:])
            for r in p.runs:r.font.size=Pt(10)
            p.paragraph_format.space_after=Pt(6)
        elif line.startswith('Narration: '):
            text=line[11:];demo.add_paragraph(text)
            if in_core:spoken+=len(text.split())
        else:demo.add_paragraph(line.replace('`',''))
    demo.core_properties.comments='New narration authored from the current report, app and saved evidence; earlier demo scripts not consulted.'
    demo.save(DEMO)
    words=lambda doc:len(' '.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells]).split())
    meta=dict(report=dict(path=str(REPORT),sha256=sha(REPORT),source_sha256=before,word_count=words(d),figures=12,changed_paragraph_prefixes=list(changes),screenshot_sha256=sha(SCREEN)),
      demo=dict(path=str(DEMO),sha256=sha(DEMO),word_count=words(demo),core_spoken_words=spoken,core_duration_seconds=355,text_source_sha256=sha(SCRIPT)),scientific_results_changed=False)
    (ROOT/'output/docx/Nebula_Current_Presentation_20260915_v2_sources.json').write_text(json.dumps(meta,indent=2),encoding='utf-8')
    print(json.dumps(meta,indent=2))

if __name__=='__main__':build()
